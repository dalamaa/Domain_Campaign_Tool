"""Human-friendly campaign exports.

This module deliberately has its own format and renderers.  The normalized
Exact App Backup remains in ``backup_export_service`` and is not changed.
"""

import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.models import (
    Campaign,
    CampaignEmailBlock,
    CampaignHistory,
    EmailAccount,
)
from app.services.expiry_service import days_until_expiry
from app.services.time_service import get_business_today


BASE_DOMAIN_CAMPAIGN_COLUMNS = [
    'Domain',
    'Expiry Date',
    'Domain Status',
    'Days Left',
    'Last Contact',
    'Days Since Last Contact',
    'Sequence Start Date',
    'Total Days Active',
    'Handled By',
    'Campaign Status',
    'Current Sequence',
    'Current Price',
    'Lifecycle',
]

BASE_EMAIL_USED_COLUMNS = ['Domain', 'Lifecycle']


def _campaign_sort_key(campaign):
    return (campaign.created_at or datetime.min, campaign.id or 0)


def _campaign_context():
    contexts = {}
    for campaign in Campaign.query.all():
        ordered = sorted(campaign.domain.campaigns, key=_campaign_sort_key)
        lifecycle = next(
            ordinal for ordinal, item in enumerate(ordered, 1)
            if item.id == campaign.id
        )
        contexts[campaign.id] = {
            'campaign': campaign,
            'domain': campaign.domain,
            'lifecycle': lifecycle,
        }
    return contexts


def _human_date(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        value = value.date()
    return value.strftime('%m/%d/%Y')


def _enum_value(value):
    return value.value if hasattr(value, 'value') else value


def _email_codes(campaign, account_orders):
    codes = list(dict.fromkeys(
        block.email_code
        for block in CampaignEmailBlock.query.filter_by(campaign_id=campaign.id).all()
    ))
    return sorted(codes, key=lambda code: (account_orders.get(code, float('inf')), code))


def build_spreadsheet_datasets():
    """Return the two human-friendly datasets as ``(name, columns, rows)``."""
    contexts = _campaign_context()
    business_today = get_business_today()
    accounts = EmailAccount.query.all()
    account_orders = {account.code: account.profile_order for account in accounts}

    campaigns = sorted(
        (context['campaign'] for context in contexts.values()),
        key=lambda campaign: (campaign.domain.domain_name, *_campaign_sort_key(campaign)),
    )
    history_by_campaign = {}
    maximum_sequence = 0
    for history in CampaignHistory.query.all():
        history_by_campaign.setdefault(history.campaign_id, []).append(history)
        if history.sequence is not None:
            maximum_sequence = max(maximum_sequence, history.sequence)
    for campaign in campaigns:
        if campaign.current_sequence is not None:
            maximum_sequence = max(maximum_sequence, campaign.current_sequence)
    for history_rows in history_by_campaign.values():
        history_rows.sort(key=lambda history: (history.sequence if history.sequence is not None else -1, history.id or 0))

    domain_columns = BASE_DOMAIN_CAMPAIGN_COLUMNS + [
        f'Email Sent #{sequence} (Price)'
        for sequence in range(1, maximum_sequence + 1)
    ]
    domain_rows = []
    for campaign in campaigns:
        domain = campaign.domain
        histories = history_by_campaign.get(campaign.id, [])
        prices = {
            history.sequence: history.price_after
            for history in histories
            if history.sequence is not None
        }
        last_contact = campaign.last_contact_date
        start_date = campaign.start_date
        domain_rows.append({
            'Domain': domain.domain_name,
            'Expiry Date': _human_date(domain.expiry_date),
            'Domain Status': domain.status or '',
            'Days Left': days_until_expiry(domain.expiry_date, business_today),
            'Last Contact': _human_date(last_contact),
            'Days Since Last Contact': (
                (business_today - last_contact).days if last_contact is not None else ''
            ),
            'Sequence Start Date': _human_date(start_date),
            'Total Days Active': (
                (business_today - start_date).days if start_date is not None else ''
            ),
            'Handled By': campaign.handled_by or '',
            'Campaign Status': _enum_value(campaign.status) or '',
            'Current Sequence': campaign.current_sequence,
            'Current Price': campaign.current_price,
            'Lifecycle': contexts[campaign.id]['lifecycle'],
            **{
                f'Email Sent #{sequence} (Price)': prices.get(sequence, '')
                for sequence in range(1, maximum_sequence + 1)
            },
        })

    maximum_email_count = 0
    email_rows = []
    for campaign in campaigns:
        codes = _email_codes(campaign, account_orders)
        maximum_email_count = max(maximum_email_count, len(codes))
        email_rows.append({
            'Domain': campaign.domain.domain_name,
            'Lifecycle': contexts[campaign.id]['lifecycle'],
            **{f'Email {index}': code for index, code in enumerate(codes, 1)},
        })
    email_columns = BASE_EMAIL_USED_COLUMNS + [
        f'Email {index}' for index in range(1, maximum_email_count + 1)
    ]
    for row in email_rows:
        for column in email_columns:
            row.setdefault(column, '')

    return [
        ('Domain Campaign', domain_columns, domain_rows),
        ('Email Used', email_columns, email_rows),
    ]


def _write_sheet(sheet, columns, rows):
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    for row in rows:
        sheet.append([row.get(column, '') for column in columns])
    for index, column in enumerate(columns, 1):
        lengths = [len(column)] + [len(str(row.get(column, ''))) for row in rows]
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(lengths) + 2, 45)


def build_spreadsheet_xlsx():
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, columns, rows in build_spreadsheet_datasets():
        _write_sheet(workbook.create_sheet(name), columns, rows)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_spreadsheet_csv_zip():
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, columns, rows in build_spreadsheet_datasets():
            filename = 'domain_campaign.csv' if name == 'Domain Campaign' else 'email_used.csv'
            csv_output = io.StringIO(newline='')
            writer = csv.DictWriter(csv_output, fieldnames=columns, lineterminator='\n')
            writer.writeheader()
            writer.writerows(rows)
            archive.writestr(filename, csv_output.getvalue().encode('utf-8'))
    output.seek(0)
    return output


def spreadsheet_export_filename(extension):
    return f'domain-campaign-spreadsheet-export-{get_business_today().isoformat()}.{extension}'
