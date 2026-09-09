import csv
import io
import json
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from app.models.models import (
    ActionType,
    Campaign,
    CampaignEmailBlock,
    CampaignHistory,
    CampaignStatus,
    Domain,
    EmailAccount,
    HistoryEmailUsed,
    Reservation,
    ReservationEmailLink,
    ReservationStatus,
    Setting,
    db,
)


def _seed_export_data():
    domain = Domain(
        domain_name='example.com',
        expiry_date=date(2030, 1, 2),
        status='SOLD',
        notes='Keep this domain note',
        created_at=datetime(2025, 1, 1, 9, 0),
    )
    older_domain = Domain(domain_name='older.example', created_at=datetime(2025, 1, 1, 8, 0))
    db.session.add_all([
        domain,
        older_domain,
        EmailAccount(code='A01', group='A', profile_order=2, enabled=True),
        EmailAccount(code='T05', group='T', profile_order=1, enabled=False),
        Setting(key='BUSINESS_TIMEZONE', value='UTC'),
        Setting(key='ADMIN_PASSWORD', value='must-not-export'),
        Setting(key='SECRET_KEY', value='must-not-export'),
    ])
    db.session.flush()
    older = Campaign(
        domain_id=domain.id,
        status=CampaignStatus.RESTING,
        current_price=80,
        current_sequence=2,
        created_at=datetime(2025, 2, 1, 9, 0),
        updated_at=datetime(2025, 2, 2, 9, 0),
    )
    current = Campaign(
        domain_id=domain.id,
        status=CampaignStatus.ACTIVE,
        start_date=date(2025, 3, 1),
        current_price=100,
        current_sequence=1,
        created_at=datetime(2025, 3, 1, 9, 0),
        updated_at=datetime(2025, 3, 2, 9, 0),
    )
    other = Campaign(
        domain_id=older_domain.id,
        status=CampaignStatus.DORMANT,
        current_price=0,
        current_sequence=0,
        created_at=datetime(2025, 4, 1, 9, 0),
    )
    db.session.add_all([older, current, other])
    db.session.flush()
    history = CampaignHistory(
        campaign_id=current.id,
        sequence=1,
        action_type=ActionType.FIRST_OUTREACH,
        action_date=None,
        price_before=0,
        price_after=100,
        sequence_before=0,
        sequence_after=1,
        notes='history note',
    )
    old_history = CampaignHistory(
        campaign_id=older.id,
        sequence=1,
        action_type=ActionType.FIRST_OUTREACH,
        action_date=datetime(2025, 2, 2, 9, 0),
    )
    db.session.add_all([history, old_history])
    db.session.flush()
    db.session.add_all([
        HistoryEmailUsed(history_id=history.id, email_code='T05'),
        CampaignEmailBlock(campaign_id=current.id, email_code='T05'),
    ])
    reservation = Reservation(
        campaign_id=current.id,
        date=date(2025, 3, 3),
        status=ReservationStatus.RESERVED,
        is_override=True,
    )
    db.session.add(reservation)
    db.session.flush()
    db.session.add(ReservationEmailLink(reservation_id=reservation.id, email_code='T05'))
    db.session.commit()
    return domain, older, current, history


def _authenticated(client):
    client.application.config.update(
        SECRET_KEY='test-session-secret',
        ADMIN_USERNAME='admin',
        ADMIN_PASSWORD='test-password',
        TESTING=True,
    )
    with client.session_transaction() as session:
        session['authenticated'] = True


def test_export_endpoints_require_authentication(app, client):
    app.config.update(
        TESTING=False,
        SECRET_KEY='test-session-secret',
        ADMIN_USERNAME='admin',
        ADMIN_PASSWORD='test-password',
    )
    assert client.get('/api/backup/export.xlsx').status_code == 401
    assert client.get('/api/backup/export.zip').status_code == 401
    app.config['TESTING'] = True


def test_settings_page_exposes_export_controls(client):
    html = client.get('/settings').get_data(as_text=True)
    assert 'id="backup-export"' in html
    assert '/api/backup/export.xlsx' in html
    assert '/api/backup/export.zip' in html
    assert 'Restore Exact App Backup' in html


def test_xlsx_export_contains_expected_data_and_safe_values(app, client):
    with app.app_context():
        _seed_export_data()
    _authenticated(client)
    response = client.get('/api/backup/export.xlsx')
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment; filename=domain-campaign-backup-' in response.headers['Content-Disposition']
    assert response.headers['Content-Disposition'].endswith('.xlsx')
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    assert workbook.sheetnames == [
        'Domains', 'Campaigns', 'Campaign History', 'History Email Used',
        'Email Accounts', 'Campaign Email Associations', 'Reservations',
        'Reservation Email Links', 'Settings', 'Manifest',
    ]
    domains = list(workbook['Domains'].values)
    assert domains[1][1:5] == ('example.com', '2030-01-02', 'SOLD', 'Keep this domain note')
    campaigns = list(workbook['Campaigns'].values)
    example_rows = [row for row in campaigns[1:] if row[1] == 'example.com']
    assert [row[2] for row in example_rows] == [1, 2]
    assert [row[5] for row in example_rows] == ['RESTING', 'ACTIVE']
    assert any(row[6] == None for row in campaigns[1:])
    history = list(workbook['Campaign History'].values)
    assert all(row[5] == 'FIRST_OUTREACH' for row in history[1:])
    assert any(row[6] in ('', None) for row in history[1:])
    assert ('must-not-export' not in response.data.decode('latin1'))
    email_rows = list(workbook['Email Accounts'].values)
    assert next(row for row in email_rows[1:] if row[0] == 'T05')[3] == 'false'


def test_csv_zip_export_has_normalized_files_and_logical_references(app, client):
    with app.app_context():
        _seed_export_data()
    _authenticated(client)
    response = client.get('/api/backup/export.zip')
    assert response.status_code == 200
    assert response.mimetype == 'application/zip'
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert archive.namelist() == [
            'domains.csv', 'campaigns.csv', 'campaign_history.csv',
            'history_email_used.csv', 'email_accounts.csv',
            'campaign_email_associations.csv', 'reservations.csv',
            'reservation_email_links.csv', 'settings.csv', 'manifest.json',
        ]
        campaigns = list(csv.DictReader(io.StringIO(archive.read('campaigns.csv').decode('utf-8'))))
        assert {row['lifecycle_ordinal'] for row in campaigns if row['domain_name'] == 'example.com'} == {'1', '2'}
        links = list(csv.DictReader(io.StringIO(archive.read('campaign_email_associations.csv').decode('utf-8'))))
        assert links[0]['domain_name'] == 'example.com'
        assert links[0]['email_code'] == 'T05'
        settings = archive.read('settings.csv').decode('utf-8')
        assert 'BUSINESS_TIMEZONE' in settings
        assert 'ADMIN_PASSWORD' not in settings
        assert 'SECRET_KEY' not in settings


def test_export_is_read_only_and_empty_database_still_has_headers(app, client):
    _authenticated(client)
    with app.app_context():
        before = [
            Domain.query.count(), Campaign.query.count(), CampaignHistory.query.count(),
            EmailAccount.query.count(), Reservation.query.count(), Setting.query.count(),
        ]
    response = client.get('/api/backup/export.zip')
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        for name in archive.namelist():
            content = archive.read(name).decode('utf-8')
            assert content.strip()
            assert content.splitlines()[0]
    with app.app_context():
        after = [
            Domain.query.count(), Campaign.query.count(), CampaignHistory.query.count(),
            EmailAccount.query.count(), Reservation.query.count(), Setting.query.count(),
        ]
    assert before == after == [0, 0, 0, 0, 0, 0]


def test_exact_zip_manifest_describes_every_csv_without_changing_csv_data(app, client):
    with app.app_context():
        _seed_export_data()
    _authenticated(client)
    response = client.get('/api/backup/export.zip')
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['format_name'] == 'domain-campaign-exact-backup'
        assert manifest['format_version'] == 1
        assert manifest['application'] == 'Domain Campaign Tool'
        datetime.fromisoformat(manifest['exported_at'])
        date.fromisoformat(manifest['business_date'])
        assert len(manifest['datasets']) == 9
        for dataset in manifest['datasets']:
            rows = list(csv.reader(io.StringIO(archive.read(dataset['filename']).decode('utf-8'))))
            assert dataset['columns'] == rows[0]
            assert dataset['row_count'] == len(rows) - 1
        assert 'manifest.json' in archive.namelist()


def test_exact_xlsx_manifest_matches_all_data_sheets(app, client):
    with app.app_context():
        _seed_export_data()
    _authenticated(client)
    workbook = load_workbook(io.BytesIO(client.get('/api/backup/export.xlsx').data), read_only=False)
    manifest = workbook['Manifest']
    metadata = {
        row[0].value: row[1].value
        for row in manifest.iter_rows(min_row=1, max_row=7)
    }
    assert metadata['format_name'] == 'domain-campaign-exact-backup'
    assert metadata['format_version'] == 1
    assert metadata['application'] == 'Domain Campaign Tool'
    datetime.fromisoformat(metadata['exported_at'])
    date.fromisoformat(metadata['business_date'])
    assert metadata['schema_revision'] in ('', None)
    assert [cell.value for cell in manifest[9]] == [
        'dataset', 'sheet_name', 'filename', 'columns', 'row_count',
    ]
    for row in manifest.iter_rows(min_row=10, max_row=18, values_only=True):
        dataset, sheet_name, _filename, columns_json, row_count = row
        sheet = workbook[sheet_name]
        assert json.loads(columns_json) == [cell.value for cell in sheet[1]]
        assert row_count == sheet.max_row - 1


def test_exact_empty_export_manifest_has_zero_dataset_counts(app, client):
    _authenticated(client)
    with zipfile.ZipFile(io.BytesIO(client.get('/api/backup/export.zip').data)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert all(dataset['row_count'] == 0 for dataset in manifest['datasets'])


def test_unavailable_optional_manifest_metadata_does_not_break_export(app, client, monkeypatch):
    with app.app_context():
        _seed_export_data()
    _authenticated(client)
    import app.services.backup_export_service as export_service
    monkeypatch.setattr(export_service, '_git_revision', lambda: None)
    monkeypatch.setattr(export_service, '_schema_revision', lambda: None)
    response = client.get('/api/backup/export.zip')
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['git_revision'] is None
        assert manifest['schema_revision'] is None
