import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from app.models.models import (
    ActionType,
    Campaign,
    CampaignEmailBlock,
    CampaignHistory,
    CampaignStatus,
    Domain,
    EmailAccount,
    Reservation,
    db,
)


def _authenticate(client):
    client.application.config.update(
        TESTING=True,
        SECRET_KEY='test-session-secret',
        ADMIN_USERNAME='admin',
        ADMIN_PASSWORD='test-password',
    )
    with client.session_transaction() as session:
        session['authenticated'] = True


def _seed_spreadsheet_data():
    domain = Domain(
        domain_name='LandscapeDesignBoston.com',
        expiry_date=date(2026, 11, 15),
        status='AVAILABLE',
        created_at=datetime(2025, 1, 1, 9, 0),
    )
    dormant_domain = Domain(domain_name='dormant.example.com', created_at=datetime(2025, 1, 2, 9, 0))
    accounts = [
        EmailAccount(code='ML06', group='ML', profile_order=2, enabled=False),
        EmailAccount(code='ML07', group='ML', profile_order=3, enabled=True),
        EmailAccount(code='N04', group='N', profile_order=1, enabled=True),
        EmailAccount(code='N09', group='N', profile_order=4, enabled=True),
    ]
    db.session.add_all([domain, dormant_domain, *accounts])
    db.session.flush()
    first = Campaign(
        domain_id=domain.id,
        status=CampaignStatus.RESTING,
        current_sequence=1,
        current_price=450,
        created_at=datetime(2025, 2, 1, 9, 0),
    )
    current = Campaign(
        domain_id=domain.id,
        status=CampaignStatus.ACTIVE,
        start_date=date(2026, 1, 1),
        last_contact_date=date(2026, 8, 10),
        handled_by='Michael',
        current_sequence=3,
        current_price=395,
        created_at=datetime(2025, 3, 1, 9, 0),
    )
    dormant = Campaign(
        domain_id=dormant_domain.id,
        status=CampaignStatus.DORMANT,
        current_sequence=0,
        current_price=0,
        created_at=datetime(2025, 4, 1, 9, 0),
    )
    db.session.add_all([first, current, dormant])
    db.session.flush()
    db.session.add_all([
        CampaignHistory(
            campaign_id=current.id,
            sequence=1,
            action_type=ActionType.FIRST_OUTREACH,
            action_date=datetime(2026, 7, 1, 9, 0),
            price_before=0,
            price_after=450,
        ),
        CampaignHistory(
            campaign_id=current.id,
            sequence=3,
            action_type=ActionType.PRICE_REDUCTION,
            action_date=datetime(2026, 8, 10, 9, 0),
            price_before=450,
            price_after=395,
        ),
    ])
    db.session.flush()
    db.session.add_all([
        CampaignEmailBlock(campaign_id=first.id, email_code='N04'),
        CampaignEmailBlock(campaign_id=current.id, email_code='ML06'),
        CampaignEmailBlock(campaign_id=current.id, email_code='ML07'),
        CampaignEmailBlock(campaign_id=current.id, email_code='N04'),
        CampaignEmailBlock(campaign_id=current.id, email_code='N09'),
    ])
    db.session.commit()


def test_all_export_endpoints_require_authentication(app, client):
    app.config.update(
        TESTING=False,
        SECRET_KEY='test-session-secret',
        ADMIN_USERNAME='admin',
        ADMIN_PASSWORD='test-password',
    )
    for endpoint in (
        '/api/backup/export.xlsx',
        '/api/backup/export.zip',
        '/api/export/spreadsheet.xlsx',
        '/api/export/spreadsheet.zip',
    ):
        assert client.get(endpoint).status_code == 401


def test_spreadsheet_xlsx_has_human_friendly_lifecycle_rows(app, client, monkeypatch):
    with app.app_context():
        _seed_spreadsheet_data()
    _authenticate(client)
    import app.services.spreadsheet_export_service as export_service
    monkeypatch.setattr(export_service, 'get_business_today', lambda: date(2026, 9, 7))

    response = client.get('/api/export/spreadsheet.xlsx')
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ['Domain Campaign', 'Email Used']
    sheet = workbook['Domain Campaign']
    assert sheet.freeze_panes == 'A2'
    rows = list(sheet.values)
    headers = rows[0]
    assert headers[:13] == (
        'Domain', 'Expiry Date', 'Domain Status', 'Days Left', 'Last Contact',
        'Days Since Last Contact', 'Sequence Start Date', 'Total Days Active',
        'Handled By', 'Campaign Status', 'Current Sequence', 'Current Price',
        'Lifecycle',
    )
    assert headers[-3:] == ('Email Sent #1 (Price)', 'Email Sent #2 (Price)', 'Email Sent #3 (Price)')
    campaign_rows = [dict(zip(headers, row)) for row in rows[1:]]
    current_rows = [row for row in campaign_rows if row['Domain'] == 'LandscapeDesignBoston.com']
    assert [row['Lifecycle'] for row in current_rows] == [1, 2]
    active = next(row for row in current_rows if row['Campaign Status'] == 'ACTIVE')
    assert active['Expiry Date'] == '11/15/2026'
    assert active['Last Contact'] == '08/10/2026'
    assert active['Current Sequence'] == 3
    assert active['Current Price'] == 395
    assert active['Email Sent #1 (Price)'] == 450
    assert active['Email Sent #2 (Price)'] in ('', None)
    assert active['Email Sent #3 (Price)'] == 395
    dormant = next(row for row in campaign_rows if row['Domain'] == 'dormant.example.com')
    assert (dormant['Campaign Status'], dormant['Current Sequence'], dormant['Current Price']) == ('DORMANT', 0, 0)

    email_sheet = list(workbook['Email Used'].values)
    email_headers = email_sheet[0]
    email_rows = [dict(zip(email_headers, row)) for row in email_sheet[1:]]
    active_emails = next(row for row in email_rows if row['Domain'] == 'LandscapeDesignBoston.com' and row['Lifecycle'] == 2)
    assert [active_emails[f'Email {index}'] for index in range(1, 5)] == ['N04', 'ML06', 'ML07', 'N09']
    assert 'ML06' in active_emails.values()  # disabled historical associations remain present


def test_spreadsheet_csv_zip_matches_logical_datasets_and_empty_export(app, client, monkeypatch):
    with app.app_context():
        _seed_spreadsheet_data()
    _authenticate(client)
    import app.services.spreadsheet_export_service as export_service
    monkeypatch.setattr(export_service, 'get_business_today', lambda: date(2026, 9, 7))
    response = client.get('/api/export/spreadsheet.zip')
    assert response.status_code == 200
    assert response.headers['Content-Disposition'].endswith('domain-campaign-spreadsheet-export-2026-09-07.zip')
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        assert archive.namelist() == ['domain_campaign.csv', 'email_used.csv']
        domain_rows = list(csv.DictReader(io.StringIO(archive.read('domain_campaign.csv').decode('utf-8'))))
        email_rows = list(csv.DictReader(io.StringIO(archive.read('email_used.csv').decode('utf-8'))))
    assert len([row for row in domain_rows if row['Domain'] == 'LandscapeDesignBoston.com']) == 2
    assert {row['Lifecycle'] for row in domain_rows if row['Domain'] == 'LandscapeDesignBoston.com'} == {'1', '2'}
    assert any(row['Email Sent #2 (Price)'] == '' for row in domain_rows)
    assert next(row for row in email_rows if row['Domain'] == 'LandscapeDesignBoston.com' and row['Lifecycle'] == '2')['Email 1'] == 'N04'

    # A fresh fixture is not needed: remove seeded rows, then verify headers-only output.
    with app.app_context():
        for model in (CampaignEmailBlock, CampaignHistory, Campaign, Domain, EmailAccount):
            model.query.delete()
        db.session.commit()
    empty = client.get('/api/export/spreadsheet.zip')
    with zipfile.ZipFile(io.BytesIO(empty.data)) as archive:
        assert archive.read('domain_campaign.csv').decode('utf-8').splitlines()[0].startswith('Domain,Expiry Date')
        assert archive.read('email_used.csv').decode('utf-8').splitlines()[0] == 'Domain,Lifecycle'


def test_spreadsheet_export_does_not_mutate_database(app, client, monkeypatch):
    with app.app_context():
        _seed_spreadsheet_data()
    _authenticate(client)
    import app.services.spreadsheet_export_service as export_service
    monkeypatch.setattr(export_service, 'get_business_today', lambda: date(2026, 9, 7))
    with app.app_context():
        before = (Domain.query.count(), Campaign.query.count(), CampaignHistory.query.count(), CampaignEmailBlock.query.count(), Reservation.query.count())
    client.get('/api/export/spreadsheet.xlsx')
    client.get('/api/export/spreadsheet.zip')
    with app.app_context():
        after = (Domain.query.count(), Campaign.query.count(), CampaignHistory.query.count(), CampaignEmailBlock.query.count(), Reservation.query.count())
    assert before == after
