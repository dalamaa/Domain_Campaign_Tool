import io
import re
import sys
import zipfile

from openpyxl import load_workbook

from app import create_app
from app.models.models import db
from tests.test_config import TestConfig


def test_backup_export_command_is_registered(app):
    result = app.test_cli_runner().invoke(args=["backup", "--help"])

    assert result.exit_code == 0
    assert "export" in result.output


def test_backup_export_command_writes_both_existing_formats_without_authentication(app, tmp_path):
    output_dir = tmp_path / "backups"
    output_dir.mkdir()

    result = app.test_cli_runner().invoke(
        args=["backup", "export", "--output-dir", str(output_dir)]
    )

    assert result.exit_code == 0, result.output
    assert "status=ok" in result.output

    exact_match = re.search(
        r"exact_zip=(.+/domain_campaign_exact_(\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2})\.zip)",
        result.output,
    )
    human_match = re.search(
        r"human_xlsx=(.+/domain_campaign_human_(\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2})\.xlsx)",
        result.output,
    )
    assert exact_match and human_match
    assert exact_match.group(2) == human_match.group(2)

    exact_path = output_dir / f"domain_campaign_exact_{exact_match.group(2)}.zip"
    human_path = output_dir / f"domain_campaign_human_{human_match.group(2)}.xlsx"
    assert exact_path.is_file()
    assert human_path.is_file()

    with zipfile.ZipFile(exact_path) as archive:
        assert "manifest.json" in archive.namelist()
        assert archive.read("manifest.json")

    workbook = load_workbook(io.BytesIO(human_path.read_bytes()), read_only=True)
    assert workbook.sheetnames == ["Domain Campaign", "Email Used"]


def test_backup_export_command_rejects_missing_output_directory(app, tmp_path):
    result = app.test_cli_runner().invoke(
        args=["backup", "export", "--output-dir", str(tmp_path / "missing")]
    )

    assert result.exit_code != 0


def test_repeated_backup_exports_use_distinct_timestamps(app, tmp_path):
    runner = app.test_cli_runner()
    args = ["backup", "export", "--output-dir", str(tmp_path)]

    first = runner.invoke(args=args)
    second = runner.invoke(args=args)

    assert first.exit_code == 0, first.output
    assert second.exit_code == 0, second.output
    exact_names = sorted(path.name for path in tmp_path.glob("domain_campaign_exact_*.zip"))
    human_names = sorted(path.name for path in tmp_path.glob("domain_campaign_human_*.xlsx"))
    assert len(exact_names) == len(human_names) == 2
    assert exact_names[0] != exact_names[1]
    assert exact_names[0].replace("domain_campaign_exact_", "").replace(".zip", "") in human_names[0]
    assert exact_names[1].replace("domain_campaign_exact_", "").replace(".zip", "") in human_names[1]


def test_backup_export_generation_failure_leaves_no_files(app, tmp_path, monkeypatch):
    import app.cli as cli

    def fail_human_export():
        raise RuntimeError("injected generation failure")

    monkeypatch.setattr(cli.spreadsheet_export_service, "build_spreadsheet_xlsx", fail_human_export)

    result = app.test_cli_runner().invoke(
        args=["backup", "export", "--output-dir", str(tmp_path)]
    )

    assert result.exit_code != 0
    assert list(tmp_path.iterdir()) == []


def test_backup_export_publish_failure_cleans_temporary_and_partial_files(app, tmp_path, monkeypatch):
    import app.cli as cli

    real_replace = cli.os.replace
    calls = {"count": 0}

    def fail_on_second_publish(source, destination):
        calls["count"] += 1
        if calls["count"] == 2:
            raise OSError("injected publish failure")
        return real_replace(source, destination)

    monkeypatch.setattr(cli.os, "replace", fail_on_second_publish)

    result = app.test_cli_runner().invoke(
        args=["backup", "export", "--output-dir", str(tmp_path)]
    )

    assert result.exit_code != 0
    assert list(tmp_path.iterdir()) == []


def test_backup_export_cli_does_not_start_scheduler(monkeypatch, tmp_path):
    import app.scheduler as scheduler

    started = []
    monkeypatch.setattr(scheduler, "init_scheduler", lambda app: started.append(app))
    monkeypatch.setenv("FLASK_RUN_FROM_CLI", "true")
    monkeypatch.setattr(
        sys,
        "argv",
        ["flask", "backup", "export", "--output-dir", str(tmp_path)],
    )

    class CliConfig(TestConfig):
        TESTING = False
        SCHEDULER_ENABLED = True

    cli_app = create_app(CliConfig)
    with cli_app.app_context():
        db.create_all()

    result = cli_app.test_cli_runner().invoke(
        args=["backup", "export", "--output-dir", str(tmp_path)]
    )

    assert result.exit_code == 0, result.output
    assert started == []

    with cli_app.app_context():
        db.drop_all()
